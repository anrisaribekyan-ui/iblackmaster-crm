"""Перенос справочника контрагентов из LiveSklad (выгрузка «Контрагенты.xlsx»). Пишет и меняет только Claude.

    python -m app.importers.livesklad_counteragents /app/ls/Контрагенты.xlsx            # пробный прогон
    python -m app.importers.livesklad_counteragents /app/ls/Контрагенты.xlsx --commit   # записать

Уже есть в CRM (по телефону, без телефона — по имени) — не дублируем, а только дополняем пустые поля
(почта, адрес, ИНН, признак поставщика, реквизиты). Нет в CRM — создаём. Балансы здесь не переносим:
долги и оплаты приходят с заказами и кассой.
"""

import argparse
import sys
import time
from collections import Counter

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.models import Counteragent
from app.models.counteragent import CounteragentType, HowKnow
from app.utils.phone import normalize_phone

EXTRA_FIELDS = {"ОГРН": "ogrn", "Директор": "director", "Договор": "contract", "Серия/номер": "passport",
                "Дата выдачи": "passport_date", "Кем выдан": "passport_issuer", "Код подразделения": "passport_code",
                "Дата рождения": "birth_date", "Место рождения": "birth_place", "Наименование банка": "bank",
                "БИК": "bik", "SWIFT": "swift", "Кор/счет": "corr_account", "Рас/счет": "account"}


def text(value) -> str | None:
    value = str(value).strip() if value is not None else ""
    return value or None


def clean_name(value) -> str:
    return (text(value) or "Без имени").lstrip(",").strip() or "Без имени"


def phones_of(raw) -> list[str]:
    result = []
    for part in str(raw or "").replace(";", ",").split(","):
        phone = normalize_phone(part)
        if phone and phone not in result:
            result.append(phone)
    return result


def run(db: Session, path: str, commit: bool) -> int:
    started = time.time()
    rows = load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    if "Имя" not in header or "Телефон" not in header:
        sys.exit("В файле нет колонок «Имя» и «Телефон» — это не выгрузка контрагентов LiveSklad?")

    by_phone: dict[str, Counteragent] = {}
    by_name: dict[str, Counteragent] = {}
    for c in db.scalars(select(Counteragent)):
        for phone in (c.phones or "").split(","):
            if phone.strip():
                by_phone.setdefault(phone.strip(), c)
        if not c.phones:
            by_name.setdefault(c.name.casefold(), c)
    types = {t.name.casefold(): t for t in db.scalars(select(CounteragentType))}
    sources = {h.name.casefold(): h for h in db.scalars(select(HowKnow))}

    stats = Counter()
    for values in rows:
        if not values or not any(values):
            continue
        row = dict(zip(header, values))
        name = clean_name(row.get("Имя"))
        phones = phones_of(row.get("Телефон"))
        found = next((by_phone[p] for p in phones if p in by_phone), None)
        if found is None and not phones:
            found = by_name.get(name.casefold())
        extra = {key: text(row.get(col)) for col, key in EXTRA_FIELDS.items() if text(row.get(col))}
        email, address, inn = text(row.get("Email")), text(row.get("Адрес")), text(row.get("ИНН"))
        vendor = row.get("Поставщик") == "Да"
        if found is not None:
            changed = False
            for attr, value in (("email", email), ("address", address), ("inn", inn)):
                if value and not getattr(found, attr):
                    setattr(found, attr, value)
                    changed = True
            if vendor and not found.is_vendor:
                found.is_vendor = True
                changed = True
            missing = {k: v for k, v in extra.items() if k not in (found.extra or {})}
            if missing:
                found.extra = {**(found.extra or {}), **missing}
                changed = True
            stats["дополнено" if changed else "уже есть без изменений"] += 1
            continue
        type_name = text(row.get("Тип контрагента"))
        ca_type = None
        if type_name:
            ca_type = types.get(type_name.casefold())
            if ca_type is None:
                ca_type = CounteragentType(name=type_name)
                db.add(ca_type)
                db.flush()
                types[type_name.casefold()] = ca_type
        source = text(row.get("Рекламный источник"))
        how = None
        if source:
            how = sources.get(source.casefold())
            if how is None:
                how = HowKnow(name=source)
                db.add(how)
                db.flush()
                sources[source.casefold()] = how
        client = Counteragent(name=name, phones=",".join(phones) or None, email=email, address=address, inn=inn,
                              type_id=ca_type.id if ca_type else None, how_know_id=how.id if how else None,
                              is_buyer=row.get("Покупатель") != "Нет", is_vendor=vendor, extra=extra)
        db.add(client)
        for phone in phones:
            by_phone[phone] = client
        if not phones:
            by_name[name.casefold()] = client
        stats["создано"] += 1
        stats["из них поставщиков"] += 1 if vendor else 0

    print("\nИтог:")
    for key, value in stats.items():
        print(f"  {key}: {value}")
    if commit:
        db.commit()
        print(f"\nЗаписано в базу за {time.time() - started:.0f} с.")
    else:
        db.rollback()
        print(f"\nПробный прогон за {time.time() - started:.0f} с — ничего не сохранено. Для записи добавьте --commit.")
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Перенос контрагентов из выгрузки LiveSklad")
    parser.add_argument("path", help="Файл «Контрагенты.xlsx»")
    parser.add_argument("--commit", action="store_true", help="Записать (без флага — пробный прогон)")
    args = parser.parse_args(argv)
    db = SessionLocal()
    try:
        return run(db, args.path, args.commit)
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
