"""Перенос кассовой истории из LiveSklad (выгрузка «Транзакции» в Excel). Пишет и меняет только Claude.

Запуск (на сервере, файл — в контейнер: docker compose cp Транзакции.xlsx backend:/app/transactions.xlsx):
    python -m app.importers.livesklad_cash /app/transactions.xlsx                 # пробный прогон
    python -m app.importers.livesklad_cash /app/transactions.xlsx --commit        # записать
    python -m app.importers.livesklad_cash /app/transactions.xlsx --until "30.10.2026 17:00" --commit

Что делаем:
- каждая строка выгрузки — одна транзакция в CRM: дата, касса, статья, нал/безнал, сумма, ответственный, примечание;
  «Оплата за заказ» и возвраты по заказу привязываются к заказу по номеру (если заказ уже перенесён);
- сами транзакции пишем напрямую, а не через сервис денег: комиссия банка и так лежит в выгрузке отдельной строкой,
  а «оплачено» по заказам перенесено вместе с заказами — второй раз их считать нельзя;
- остатки касс = то, что уже было в CRM, + сумма перенесённых операций. Для пустой CRM это остаток LiveSklad;
  в конце сверяем с колонкой «Остаток» последней операции каждой кассы и пишем расхождения;
- суммы округляем до рубля (в LiveSklad бывают копейки, в CRM — целые рубли), расхождение выводим в отчёте.

Защита от двойного ввода: если в CRM уже есть операции не позже последней операции файла — отказ.
Чтобы всё же догрузить (например, LiveSklad вели параллельно), укажите --until: берутся строки файла до этого
момента (по Москве), а в CRM не должно быть операций раньше него.
"""

import argparse
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.importers.livesklad_orders import MSK, parse_dt
from app.models import CashItem, CashRegister, Employee, Location, Order, Transaction

BANK_FORMS = {"безналом", "по карте", "безналичные", "карта", "перевод"}
NOTE_PREFIX = "LiveSklad"


def money(value) -> Decimal:
    return Decimal(str(value or 0).replace(" ", "").replace(",", ".").replace("–", "-")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)


def exact(value) -> Decimal:
    return Decimal(str(value or 0).replace(" ", "").replace(",", ".").replace("–", "-"))


def find_employee(people: list[Employee], text: str) -> int | None:
    """Сотрудник по «Получатель» из LiveSklad: совпадение по short_name, «фамилия имя» или по фамилии (если однозначно)."""
    text = " ".join(text.lower().split())
    if not text:
        return None
    exactly = [e for e in people if text in {(e.short_name or "").lower(), f"{e.surname or ''} {e.name}".strip().lower(),
                                              f"{e.name} {e.surname or ''}".strip().lower()}]
    if len(exactly) == 1:
        return exactly[0].id
    parts = [w for w in re.split(r"[\s.]+", text) if len(w) > 2]
    found = [e for e in people if any(w in f"{e.surname or ''} {e.name} {e.short_name or ''}".lower() for w in parts)]
    return found[0].id if len(found) == 1 else None


def read_rows(path: str):
    workbook = load_workbook(path, read_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    for required in ("Дата", "Касса", "Статья", "Сумма"):
        if required not in header:
            sys.exit(f"В файле нет колонки «{required}» — это не выгрузка транзакций LiveSklad?")
    for values in rows:
        if values and values[0] is not None or (values and any(values)):
            yield dict(zip(header, values))


def run(db: Session, path: str, commit: bool, until: datetime | None, targets: list[str] | None = None) -> int:
    started = time.time()
    registers = {r.name: r for r in db.scalars(select(CashRegister))}
    locations = {l.name: l for l in db.scalars(select(Location))}
    items = {(i.name, i.is_income): i for i in db.scalars(select(CashItem))}
    orders = {n: i for n, i in db.execute(select(Order.number, Order.id))}
    employees = {e.name: e.id for e in db.scalars(select(Employee))}
    owner = db.scalars(select(Employee).where(Employee.is_owner.is_(True))).first()
    people = list(db.scalars(select(Employee)))

    rows = []
    for row in read_rows(path):
        date = parse_dt(row.get("Дата"))
        if date is None:
            continue
        if until is not None and date > until:
            continue
        rows.append((date, row))
    if not rows:
        print("В файле нет операций для переноса.")
        return 1
    last_date = max(d for d, _ in rows)
    existing = db.scalar(select(func.count()).select_from(Transaction).where(Transaction.date <= last_date)) or 0
    if existing:
        print(f"В CRM уже есть операции кассы ({existing} шт.) не позже {last_date.astimezone(MSK):%d.%m.%Y %H:%M}.")
        print("Повторно не ввожу, чтобы не задвоить деньги. Если LiveSklad вели параллельно с CRM — укажите --until \"дд.мм.гггг чч:мм\"")
        print("(момент, с которого вы работаете только в CRM), и в CRM не должно быть операций раньше него.")
        return 1

    problems: list[str] = []
    stats = Counter()
    delta: dict[tuple[int, bool], Decimal] = defaultdict(Decimal)
    expected: dict[tuple[int, bool], tuple[datetime, Decimal]] = {}
    created_items: list[str] = []
    rounding_drift = Decimal("0")
    batch: list[Transaction] = []

    # Старые операции первыми: так остатки и порядок id совпадают с хронологией
    for date, row in sorted(rows, key=lambda x: x[0]):
        register = registers.get(str(row.get("Касса") or "").strip())
        if register is None:
            problems.append(f"нет кассы «{row.get('Касса')}» (операция {row.get('Номер') or '-'} от {date.astimezone(MSK):%d.%m.%Y})")
            stats["пропущено: нет кассы"] += 1
            continue
        raw_amount = exact(row.get("Сумма"))
        amount = money(raw_amount)
        if amount == 0:
            stats["пропущено: нулевая сумма"] += 1
            continue
        is_income = amount > 0
        name = str(row.get("Статья") or "Прочее").strip()
        item = items.get((name, is_income))
        if item is None:
            item = CashItem(name=name, is_income=is_income, type=None, affects_balance=False)
            db.add(item)
            db.flush()
            items[(name, is_income)] = item
            created_items.append(f"{name} ({'приход' if is_income else 'расход'})")
        form = str(row.get("Форма оплаты") or "").strip().lower()
        is_bank = form in BANK_FORMS or (not register.accepts_cash and register.accepts_bank)
        number = str(row.get("Номер") or "").strip()
        order_id = orders.get(number) if item.type in ("order", "orderReturn") else None
        location = locations.get(str(row.get("Локация") or "").strip())
        responsible = employees.get(str(row.get("Ответственный") or "").strip()) or (owner.id if owner else None)
        note_parts = [p for p in (str(row.get("Примечание") or "").strip(),) if p]
        if number and order_id is None:
            note_parts.append(f"док. {number}")
        recipient = str(row.get("Получатель") or "").strip()
        if recipient:
            note_parts.append(f"получатель: {recipient}")
        if row.get("Номер ОФД"):
            note_parts.append(f"ОФД {row['Номер ОФД']}")
        note_parts.append(f"{NOTE_PREFIX}")
        rest = row.get("Остаток")
        pay_to = find_employee(people, recipient) if item.type == "salary" else None
        if item.type == "salary" and recipient and pay_to is None:
            problems.append(f"выплата зарплаты: не нашёл сотрудника «{recipient}» ({date.astimezone(MSK):%d.%m.%Y})")
        tx = Transaction(
            date=date,
            cash_register_id=register.id,
            cash_item_id=item.id,
            is_income=is_income,
            is_bank=is_bank,
            amount=abs(amount),
            balance_after=money(rest) if rest not in (None, "") else Decimal("0"),
            location_id=location.id if location else register.location_id,
            created_by_id=responsible,
            employee_id=pay_to,
            order_id=order_id,
            note="; ".join(note_parts)[:1000],
        )
        batch.append(tx)
        delta[(register.id, is_bank)] += amount
        rounding_drift += amount - raw_amount
        if rest not in (None, ""):
            expected[(register.id, is_bank)] = (date, exact(rest))
        stats["операций"] += 1
        stats["привязано к заказу"] += 1 if order_id else 0
        if len(batch) >= 2000:
            db.add_all(batch)
            db.flush()
            batch = []
    db.add_all(batch)
    db.flush()

    names = {r.id: r.name for r in registers.values()}
    print("\nИтог:")
    for key, value in stats.items():
        print(f"  {key}: {value}")
    if created_items:
        print(f"  созданы статьи денег: {', '.join(sorted(set(created_items)))}")
    print(f"  округление копеек до рубля, суммарно: {rounding_drift:+} ₽")
    print("\nОстатки касс после переноса (CRM) и остаток по LiveSklad на последней операции:")
    for (register_id, is_bank), shift in sorted(delta.items(), key=lambda x: names[x[0][0]]):
        register = db.get(CashRegister, register_id)
        if is_bank:
            register.bank_balance += shift
            now = register.bank_balance
        else:
            register.cash_balance += shift
            now = register.cash_balance
        ls = expected.get((register_id, is_bank))
        mark = ""
        if ls is not None and abs(now - ls[1]) > Decimal("5"):
            mark = f"   ← расхождение {now - ls[1]:+}"
        print(f"  {names[register_id]:28} {'безнал' if is_bank else 'нал':7} {now:>16,.2f}   LiveSklad: {ls[1] if ls else '—':>16}{mark}")
    # Приводим остатки к снятым в LiveSklad: разницу фиксируем отдельной операцией, чтобы история и итог сошлись
    for spec in targets or []:
        try:
            name, kind, value = spec.split(":")
            target = Decimal(value.replace(",", ".")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        except ValueError:
            problems.append(f"не понял --set-balance «{spec}», формат «Касса Лента:нал:0»")
            continue
        register = registers.get(name.strip())
        if register is None:
            problems.append(f"--set-balance: нет кассы «{name}»")
            continue
        is_bank = kind.strip().lower().startswith("без")
        current = register.bank_balance if is_bank else register.cash_balance
        diff = target - current
        if diff == 0:
            continue
        label = "Корректировка остатка (LiveSklad)"
        item = items.get((label, diff > 0))
        if item is None:
            item = CashItem(name=label, is_income=diff > 0, type=None, affects_balance=False)
            db.add(item)
            db.flush()
            items[(label, diff > 0)] = item
        db.add(Transaction(date=last_date, cash_register_id=register.id, cash_item_id=item.id, is_income=diff > 0,
                           is_bank=is_bank, amount=abs(diff), balance_after=target, location_id=register.location_id,
                           created_by_id=owner.id if owner else None,
                           note=f"{NOTE_PREFIX}: приведение остатка к LiveSklad ({current} → {target})"))
        if is_bank:
            register.bank_balance = target
        else:
            register.cash_balance = target
        print(f"  корректировка {register.name} ({'безнал' if is_bank else 'нал'}): {diff:+,} → {target:,}")
    for problem in problems[:20]:
        print("  !", problem)
    if len(problems) > 20:
        print(f"  … и ещё {len(problems) - 20}")

    if commit:
        db.commit()
        print(f"\nЗаписано в базу за {time.time() - started:.0f} с.")
    else:
        db.rollback()
        print(f"\nПробный прогон за {time.time() - started:.0f} с — ничего не сохранено. Для записи добавьте --commit.")
    return 1 if problems else 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Перенос кассовой истории из выгрузки LiveSklad")
    parser.add_argument("path", help="Файл «Транзакции.xlsx» из LiveSklad")
    parser.add_argument("--commit", action="store_true", help="Записать в базу (без флага — пробный прогон)")
    parser.add_argument("--until", help="Брать операции не позже этого момента по Москве, «дд.мм.гггг чч:мм»")
    parser.add_argument("--set-balance", action="append", help="«Касса:нал|безнал:сумма» — привести остаток к LiveSklad (можно несколько раз)")
    args = parser.parse_args(argv)
    until = parse_dt(args.until) if args.until else None
    if args.until and until is None:
        sys.exit("Не понял --until. Формат: «30.10.2026 17:00»")
    db = SessionLocal()
    try:
        return run(db, args.path, args.commit, until, args.set_balance)
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
