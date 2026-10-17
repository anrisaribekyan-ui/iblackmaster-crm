"""Перенос остатков склада из LiveSklad (выгрузка «Остатки» в Excel). Пишет и меняет только Claude.

Запуск:
    python -m app.importers.livesklad_stock /path/Остатки.xlsx                          # покажет склады
    python -m app.importers.livesklad_stock /path/Остатки.xlsx --store "Панфиловский"   # пробный прогон
    python -m app.importers.livesklad_stock /path/Остатки.xlsx --store "Панфиловский" --commit

Одна выгрузка = один склад (в файле LiveSklad склада нет — его указывают в --store).

Что делаем:
- товары: находим по коду LiveSklad (код переносится как есть), нет — создаём; обновляем название, группу,
  закупочную цену, гарантию, неснижаемый остаток, артикул и цены (розничная, ремонтная, минимальная);
- группы: путь «/Магазин/Чехлы/Чехлы iPhone» раскладываем по дереву, недостающие создаём;
- остатки: одним проведённым документом «Поступление» с пометкой «Ввод остатков из LiveSklad»,
  без поставщика и без оплаты — касса не меняется. Цена единицы = себестоимость остатков / количество,
  так средняя закупочная совпадает с LiveSklad.

Защита от двойного ввода: если на складе уже есть остатки или ввод из LiveSklad уже делали — отказ.
Штрихкоды не переносим: у LiveSklad это «2000000» + код, своих штрихкодов в файле нет.
"""

import argparse
import sys
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.models import (
    Employee,
    Measure,
    Nomenclature,
    NomenclatureGroup,
    NomenclaturePrice,
    PriceType,
    StockBalance,
    StockDocument,
    Store,
)
from app.models.stock import StockDocType
from app.services import stock_documents

NOTE = "Ввод остатков из LiveSklad"
PRICE_COLUMNS = ("Розничная", "Ремонтная", "Минимальная")


def _dec(value) -> Decimal | None:
    if value in (None, ""):
        return None
    return Decimal(str(value))


def read_rows(path: str) -> list[dict]:
    import os

    if not os.path.exists(path):
        sys.exit(f"Файл «{path}» не найден. Сначала скопируйте его в контейнер: docker compose cp <файл> backend:/app/stock.xlsx")
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(rows)]
    for required in ("Наименование", "Остаток", "Код", "Себестоимость остатков"):
        if required not in header:
            sys.exit(f"В файле нет колонки «{required}» — это не выгрузка остатков LiveSklad?")
    return [dict(zip(header, r)) for r in rows if r and r[header.index("Наименование")]]


def find_group(db: Session, path: str | None, cache: dict, created: list) -> int | None:
    if not path or not str(path).strip("/ "):
        return None
    parent_id = None
    for name in [p.strip() for p in str(path).split("/") if p.strip()]:
        key = (parent_id, name.lower())
        if key not in cache:
            group = db.scalars(
                select(NomenclatureGroup).where(NomenclatureGroup.parent_id.is_(parent_id) if parent_id is None else NomenclatureGroup.parent_id == parent_id)
            ).all()
            match = next((g for g in group if g.name.strip().lower() == name.lower()), None)
            if match is None:
                match = NomenclatureGroup(name=name, parent_id=parent_id)
                db.add(match)
                db.flush()
                created.append(str(path))
            cache[key] = match.id
        parent_id = cache[key]
    return parent_id


def pick_store(db: Session, query: str | None) -> Store | None:
    stores = db.scalars(select(Store).where(Store.is_active.is_(True)).order_by(Store.id)).all()
    if not query:
        print("Укажите склад: --store \"часть названия склада или точки\". Склады:")
        for s in stores:
            print(f"  {s.id}: {s.name} — {s.location.name}")
        return None
    q = query.strip().lower()
    by_id = [s for s in stores if q == str(s.id)]  # «1» — это склад №1, а не «1000 мелочей»
    matches = by_id or [s for s in stores if q in s.name.lower() or q in s.location.name.lower()]
    if len(matches) != 1:
        print(f"По «{query}» найдено складов: {len(matches)}. Уточните. Склады:")
        for s in stores:
            print(f"  {s.id}: {s.name} — {s.location.name}")
        return None
    return matches[0]


def run(db: Session, path: str, store_query: str | None, commit: bool) -> int:
    rows = read_rows(path)
    store = pick_store(db, store_query)
    if store is None:
        return 2

    already = db.scalar(
        select(func.count(StockDocument.id)).where(
            StockDocument.store_id == store.id, StockDocument.note == NOTE, StockDocument.is_deleted.is_(False)
        )
    )
    if already:
        print(f"На склад «{store.name}» остатки из LiveSklad уже вводили. Повторно не ввожу, чтобы не задвоить.")
        return 3
    existing = db.scalar(select(func.count(StockBalance.id)).where(StockBalance.store_id == store.id, StockBalance.quantity != 0))
    if existing:
        print(f"На складе «{store.name}» уже есть остатки ({existing} позиций). Ввод из LiveSklad — только на пустой склад.")
        return 3

    owner = db.scalars(select(Employee).where(Employee.is_owner.is_(True)).order_by(Employee.id)).first()
    measure = db.scalars(select(Measure).order_by(Measure.id)).first()
    price_types = {p.name.strip().lower(): p.id for p in db.scalars(select(PriceType))}
    group_cache: dict = {}
    new_groups: list[str] = []
    created = updated = 0
    positions = []
    total_qty = Decimal("0")
    total_cost = Decimal("0")

    for row in rows:
        code = int(row["Код"])
        item = db.scalars(select(Nomenclature).where(Nomenclature.code == code)).first()
        if item is None:
            item = Nomenclature(code=code, name="", is_work=False)
            db.add(item)
            created += 1
        else:
            updated += 1
        item.name = str(row["Наименование"]).strip()[:300]
        item.is_deleted = False
        item.group_id = find_group(db, row.get("Полная группа"), group_cache, new_groups) or item.group_id
        item.measure_id = item.measure_id or (measure.id if measure else None)
        item.purchase_price = _dec(row.get("Закупочная")) or Decimal("0")
        days = _dec(row.get("Гарантия в днях"))
        months = _dec(row.get("Гарантия в месяцах"))
        item.guarantee_days = int(days) if days else int(months * 30) if months else item.guarantee_days or 0
        item.min_count = _dec(row.get("Неснижаемый остаток")) or Decimal("0")
        if row.get("Артикул"):
            item.article = str(row["Артикул"]).strip()[:100]
        percent = _dec(row.get("Зарплата, %"))
        if percent is not None:
            item.salary_percent = percent
        db.flush()

        for column in PRICE_COLUMNS:
            value = _dec(row.get(column))
            type_id = price_types.get(column.lower())
            if value is None or type_id is None:
                continue
            price = db.scalars(
                select(NomenclaturePrice).where(
                    NomenclaturePrice.nomenclature_id == item.id,
                    NomenclaturePrice.price_type_id == type_id,
                    NomenclaturePrice.location_id.is_(None),
                )
            ).first()
            if price is None:
                db.add(NomenclaturePrice(nomenclature_id=item.id, price_type_id=type_id, location_id=None, price=value))
            else:
                price.price = value

        qty = _dec(row["Остаток"]) or Decimal("0")
        if qty > 0:
            cost = _dec(row["Себестоимость остатков"]) or Decimal("0")
            unit = (cost / qty).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            positions.append({"nomenclature_id": item.id, "quantity": qty, "price": unit})
            total_qty += qty
            total_cost += cost

    document = None
    if positions:
        document = stock_documents.create_document(
            db,
            doc_type=StockDocType.PURCHASE,
            location_id=store.location_id,
            store_id=store.id,
            to_store_id=None,
            counteragent_id=None,
            responsible_id=owner.id,
            note=NOTE,
            positions=positions,
        )

    print(f"Склад: {store.name} — {store.location.name}")
    print(f"Строк в файле: {len(rows)}; товаров новых: {created}, обновлено: {updated}")
    print(f"Групп создано: {len(set(new_groups))}")
    print(f"Позиций с остатком: {len(positions)}, штук: {total_qty.normalize()}, себестоимость: {total_cost.quantize(Decimal('1'))} ₽")
    if commit:
        db.commit()
        print(f"Записано. Документ поступления № {document.number if document else '—'}")
    else:
        db.rollback()
        print("Пробный прогон — ничего не сохранено. Для записи добавьте --commit")
    return 0


def main() -> None:
    parser = argparse.ArgumentParser(description="Перенос остатков склада из LiveSklad")
    parser.add_argument("path")
    parser.add_argument("--store", help="склад: id или часть названия склада/точки")
    parser.add_argument("--commit", action="store_true")
    args = parser.parse_args()
    with SessionLocal() as db:
        sys.exit(run(db, args.path, args.store, args.commit))


if __name__ == "__main__":
    main()
