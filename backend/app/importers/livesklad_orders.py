"""Перенос заказов из LiveSklad (выгрузка «Заказы» в Excel). Пишет и меняет только Claude.

Запуск (на боевом сервере в России — в файле персональные данные клиентов):
    python -m app.importers.livesklad_orders /path/Заказы.xlsx            # пробный прогон, ничего не сохраняет
    python -m app.importers.livesklad_orders /path/Заказы.xlsx --commit   # записать в базу

Что переносится: клиенты (склейка по телефону), заказы со статусами, датами, мастером/менеджером,
устройством, неисправностью, работами и запчастями (цены и себестоимость), суммой и оплачено.
Повторный запуск безопасен: заказ с тем же номером обновляется, а не дублируется
(так перед переключением можно догрузить свежую выгрузку).

Чего НЕ делаем намеренно:
- не создаём транзакции касс за прошлые оплаты (остатки касс заводятся отдельно на дату перехода);
- не трогаем склад (запчасти давно списаны в LiveSklad, остатки переносятся отдельно);
- не начисляем зарплату за прошлые годы (история зарплаты остаётся в LiveSklad);
- паспортные и банковские поля клиентов не переносим — они не заполнены и не нужны.
"""

import argparse
import re
import secrets
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import SessionLocal
from app.models import (
    Company,
    Counteragent,
    CounteragentType,
    Employee,
    HowKnow,
    Location,
    Order,
    OrderHistory,
    OrderPosition,
    OrderStatus,
    OrderType,
    Role,
    StatusGroup,
)
from app.security import hash_password
from app.utils.phone import normalize_phone

MSK = ZoneInfo("Europe/Moscow")
POSITION_RE = re.compile(r"^\s*([\d.,]+)\s*x\s*(.+?)\s*\(([\d\s.,]+)\s*руб\.?\)\s*;?\s*$")

# Опечатки и варианты имён сотрудников в LiveSklad → одно имя
NAME_ALIASES = {"Николайй Липатунин": "Николай Липатунин", "Anri Sar.": "Sar. Anri"}
# Нынешняя команда: полное имя в LiveSklad → имя сотрудника в сиде (тогда переименуем его в полное)
TEAM = {
    "Sar. Anri": "__owner__",
    "Сарибекян Гор": "Гор",
    "Алексей Яковлев": "Алексей",
    "Евгений Черкасов": "Евгений",
    "Николай Липатунин": "Николай",
}
# Статусы, которых может не быть в справочнике: в какую группу их класть
STATUS_GROUP_GUESS = {
    "закрыт": StatusGroup.CLOSED,
    "выдан": StatusGroup.CLOSED,
    "отказ": StatusGroup.CLOSED,
    "без ремонта": StatusGroup.CLOSED,
    "готов": StatusGroup.FINISH,
    "ждет": StatusGroup.WAIT,
    "оповест": StatusGroup.WAIT,
    "согласова": StatusGroup.WAIT,
    "принят": StatusGroup.NEW,
}


def _id(item) -> int | None:
    return item.id if item is not None else None


def money(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value).replace(" ", "").replace(",", ".")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)


def parse_dt(value) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        local = value.replace(tzinfo=MSK) if value.tzinfo is None else value
    else:
        text = str(value).strip()
        for fmt in ("%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y"):
            try:
                local = datetime.strptime(text, fmt).replace(tzinfo=MSK)
                break
            except ValueError:
                continue
        else:
            return None
    return local.astimezone(timezone.utc)


def split_list(value) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in re.split(r"[,;\n]", str(value)) if part.strip()]


def text(value) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def parse_positions(value) -> list[tuple[Decimal, str, Decimal]]:
    """«1 x Замена дисплея iPhone XR (4700 руб.);» → [(1, «Замена дисплея iPhone XR», 4700)]"""
    result = []
    for line in str(value or "").split("\n"):
        if not line.strip():
            continue
        match = POSITION_RE.match(line)
        if not match:
            raise ValueError(f"Не разобрал строку позиции: {line!r}")
        qty = Decimal(match.group(1).replace(",", "."))
        result.append((qty, match.group(2).strip(), money(match.group(3))))
    return result


def spread_cost(lines: list[tuple[Decimal, str, Decimal]], cost_total: Decimal) -> list[Decimal]:
    """Себестоимость в LiveSklad-выгрузке есть только суммой по работам/запчастям — делим пропорционально цене."""
    if not lines:
        return []
    price_total = sum((qty * price for qty, _, price in lines), Decimal("0"))
    shares = []
    for qty, _, price in lines:
        part = (cost_total * qty * price / price_total) if price_total > 0 else cost_total / len(lines)
        shares.append(part.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    shares[-1] += cost_total - sum(shares)  # копеечный остаток — на последнюю строку
    return [(share / qty).quantize(Decimal("0.01")) if qty else share for share, (qty, _, _) in zip(shares, lines)]


class Importer:
    def __init__(self, db: Session):
        self.db = db
        self.stats = Counter()
        self.warnings: list[str] = []
        self.owner = db.scalars(select(Employee).where(Employee.is_owner.is_(True))).first()
        if self.owner is None:
            raise SystemExit("Нет владельца в базе — сначала запустите сид: python -m app.seed")
        self.locations = {loc.name: loc for loc in db.scalars(select(Location))}
        self.statuses = {s.name.strip().casefold(): s for s in db.scalars(select(OrderStatus))}
        self.types = {t.name.strip().casefold(): t for t in db.scalars(select(OrderType))}
        self.how_knows = {h.name.strip().casefold(): h for h in db.scalars(select(HowKnow))}
        self.ca_types = {t.name.strip().casefold(): t for t in db.scalars(select(CounteragentType))}
        self.employees: dict[str, Employee | None] = {}
        self.clients_by_phone: dict[str, Counteragent] = {}
        for client in db.scalars(select(Counteragent)):
            for phone in (client.phones or "").split(","):
                if phone.strip():
                    self.clients_by_phone.setdefault(phone.strip(), client)
        self.clients_by_name: dict[str, Counteragent] = {
            c.name.casefold(): c for c in db.scalars(select(Counteragent).where(Counteragent.phones.is_(None)))
        }
        self.orders = {o.number: o for o in db.scalars(select(Order))}
        self.inactive_role = None

    # --- справочники ---

    def location(self, name: str) -> Location:
        name = (name or "").strip()
        if name in self.locations:
            return self.locations[name]
        for key, loc in self.locations.items():  # «iBlackMaster (ТК Панфиловский)» ↔ «ТК Панфиловский»
            if key.casefold() in name.casefold() or name.casefold() in key.casefold():
                return loc
        raise SystemExit(f"Локация «{name}» не найдена — создайте её в Настройках и запустите снова")

    def status(self, name: str) -> OrderStatus:
        key = (name or "").strip().casefold()
        if key in self.statuses:
            return self.statuses[key]
        group = next((g for word, g in STATUS_GROUP_GUESS.items() if word in key), StatusGroup.IN_WORK)
        status = OrderStatus(name=name.strip(), group=group, color="#98a2b3", sort=99)
        self.db.add(status)
        self.db.flush()
        self.statuses[key] = status
        self.warnings.append(f"Создан статус «{status.name}» в группе {group.value}")
        return status

    def order_type(self, name: str) -> OrderType:
        key = (name or "Негарантийный").strip().casefold()
        if key not in self.types:
            item = OrderType(name=name.strip())
            self.db.add(item)
            self.db.flush()
            self.types[key] = item
            self.warnings.append(f"Создан тип заказа «{item.name}»")
        return self.types[key]

    def how_know(self, name) -> HowKnow | None:
        name = text(name)
        if not name:
            return None
        key = name.casefold()
        if key not in self.how_knows:
            item = HowKnow(name=name)
            self.db.add(item)
            self.db.flush()
            self.how_knows[key] = item
        return self.how_knows[key]

    def employee(self, raw) -> Employee | None:
        name = text(raw)
        if not name:
            return None
        name = NAME_ALIASES.get(name, name)
        if name in self.employees:
            return self.employees[name]
        found = self.db.scalars(select(Employee).where(Employee.name == name)).first()
        if found is None and name in TEAM:
            target = TEAM[name]
            if target == "__owner__":
                found = self.owner
            else:
                found = self.db.scalars(select(Employee).where(Employee.name == target)).first()
                if found is not None:
                    found.name = name  # «Алексей» → «Алексей Яковлев», как в LiveSklad
                    self.warnings.append(f"Сотрудник «{target}» переименован в «{name}»")
        if found is None:
            # Бывший сотрудник: нужен для истории заказов, войти не может
            if self.inactive_role is None:
                self.inactive_role = self.db.scalars(select(Role).where(Role.name == "Бывшие сотрудники")).first()
                if self.inactive_role is None:
                    self.inactive_role = Role(name="Бывшие сотрудники", permissions=[], scopes={})
                    self.db.add(self.inactive_role)
                    self.db.flush()
            slug = secrets.token_hex(4)
            found = Employee(
                name=name,
                short_name=name,
                email=f"former-{slug}@iblackmaster.local",
                password_hash=hash_password(secrets.token_urlsafe(24), rounds=4),
                role_id=self.inactive_role.id,
                is_active=False,
            )
            self.db.add(found)
            self.db.flush()
            self.warnings.append(f"Добавлен бывший сотрудник «{name}» (без входа)")
        self.employees[name] = found
        return found

    def client(self, row: dict) -> Counteragent:
        phone = normalize_phone(str(row.get("Телефон") or ""))
        name = text(row.get("Имя")) or "Без имени"
        if phone and phone in self.clients_by_phone:
            self.stats["клиент найден по телефону"] += 1
            return self.clients_by_phone[phone]
        if not phone and name.casefold() in self.clients_by_name:
            return self.clients_by_name[name.casefold()]
        ca_type = self.ca_types.get((text(row.get("Тип контрагента")) or "").casefold())
        client = Counteragent(
            name=name,
            phones=phone,
            email=text(row.get("Email")),
            address=text(row.get("Адрес")),
            note=text(row.get("Примечание контрагента")),
            type_id=ca_type.id if ca_type else None,
            how_know_id=_id(self.how_know(row.get("Источник рекламы контрагента"))),
            is_buyer=True,
            is_vendor=row.get("Поставщик") == "Да",
        )
        self.db.add(client)
        self.db.flush()
        if phone:
            self.clients_by_phone[phone] = client
        else:
            self.clients_by_name[name.casefold()] = client
        self.stats["клиентов создано"] += 1
        return client

    # --- заказ ---

    def import_row(self, row: dict) -> None:
        number = text(row.get("Номер заказа"))
        if not number:
            return
        status = self.status(row["Статус"])
        created = parse_dt(row.get("Дата создания")) or datetime.now(timezone.utc)
        finished = parse_dt(row.get("Дата готовности"))
        closed = parse_dt(row.get("Дата выдачи"))
        if status.group == StatusGroup.CLOSED:
            closed = closed or finished or created
        else:
            closed = None
        master = self.employee(row.get("Мастер"))
        manager = self.employee(row.get("Менеджер"))
        creator = self.employee(row.get("Менеджер создавший заказ")) or self.owner
        closer = self.employee(row.get("Менеджер закрывший заказ")) if closed else None

        client = self.client(row)  # до создания заказа: client() пишет в базу
        how_know_id = _id(self.how_know(row.get("Источник рекламы заказа")))
        location_id = self.location(row["Локация"]).id
        order_type_id = self.order_type(row.get("Тип заказа") or "Негарантийный").id

        order = self.orders.get(number)
        is_new = order is None
        if is_new:
            order = Order(number=number)
        order.location_id = location_id
        order.order_type_id = order_type_id
        order.status_id = status.id
        order.counteragent_id = client.id
        order.device_type = text(row.get("Тип устройства"))
        order.brand = text(row.get("Марка"))
        order.model = text(row.get("Модель"))
        order.serial = text(row.get("Серийный номер / IMEI"))
        order.color = text(row.get("Цвет"))
        order.problems = [text(row.get("Неисправность"))] if row.get("Неисправность") else []
        order.appearance = [text(row.get("Внешний вид"))] if row.get("Внешний вид") else []
        order.complete_set = split_list(row.get("Комплектация"))
        order.device_password = text(row.get("Пароль"))
        order.note = text(row.get("Комментарий приемщика"))
        order.verdict = text(row.get("Вердикт"))
        order.approximate_price = text(row.get("Ориентировочная цена"))
        order.deadline = parse_dt(row.get("Крайний срок"))
        order.is_urgent = row.get("Срочно") == "Да"
        order.how_know_id = how_know_id
        order.master_id = master.id if master else None
        order.manager_id = manager.id if manager else None
        order.created_by_id = creator.id
        order.closed_by_id = closer.id if closer else None
        order.created_at = created
        order.finished_at = finished if status.group in (StatusGroup.FINISH, StatusGroup.CLOSED) else None
        order.closed_at = closed
        order.last_action_at = closed or finished or created
        order.is_deleted = False
        order.custom_fields = order.custom_fields or {}
        if is_new:
            self.db.add(order)
            self.orders[number] = order

        # Позиции: при повторном импорте пересобираем
        works = parse_positions(row.get("Выполненные работы"))
        parts = parse_positions(row.get("Установленные запчасти"))
        work_costs = spread_cost(works, money(row.get("Себестоимость работ")))
        part_costs = spread_cost(parts, money(row.get("Себестоимость запчастей")))
        if not is_new:
            order.positions.clear()
        performer = master or manager or creator
        for is_work, lines, costs in ((True, works, work_costs), (False, parts, part_costs)):
            for (qty, name, price), cost in zip(lines, costs):
                order.positions.append(
                    OrderPosition(
                        is_work=is_work,
                        name=name[:300],
                        quantity=qty,
                        price=price,
                        sold_price=price,
                        purchase_price=cost,
                        performer_id=performer.id,
                        created_at=created,
                    )
                )
        gross = sum((qty * price for qty, _, price in works + parts), Decimal("0"))
        total = money(row.get("Сумма заказа"))
        order.discount_percent = Decimal("0")
        order.discount_sum = max(Decimal("0"), gross - total)
        order.total_price = total
        order.total_purchase = money(row.get("Себестоимость заказа"))
        order.paid = money(row.get("Оплачено по заказу"))
        self.db.flush()

        if is_new:
            self.db.add(
                OrderHistory(
                    order_id=order.id,
                    created_at=created,
                    employee_id=creator.id,
                    type="created",
                    text="Перенесён из LiveSklad",
                    data={},
                )
            )
            self.stats["заказов создано"] += 1
        else:
            self.stats["заказов обновлено"] += 1

    def finish(self) -> None:
        """Сквозная нумерация продолжается после последнего номера из LiveSklad."""
        company = self.db.scalars(select(Company)).first()
        prefix = company.order_number_prefix
        numbers = [int(n[len(prefix):]) for n in self.orders if n.startswith(prefix) and n[len(prefix):].isdigit()]
        if numbers and company.next_order_number <= max(numbers):
            company.next_order_number = max(numbers) + 1
            self.warnings.append(f"Следующий номер заказа: {prefix}{company.next_order_number}")


def read_rows(path: str):
    workbook = load_workbook(path, read_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    for values in rows:
        if values and values[0] is not None:
            yield dict(zip(header, values))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Перенос заказов из выгрузки LiveSklad")
    parser.add_argument("path", help="Файл «Заказы.xlsx» из LiveSklad")
    parser.add_argument("--commit", action="store_true", help="Записать в базу (без флага — пробный прогон)")
    args = parser.parse_args(argv)

    started = time.time()
    db = SessionLocal()
    importer = Importer(db)
    errors = []
    for index, row in enumerate(read_rows(args.path), start=2):
        try:
            with db.begin_nested():
                importer.import_row(row)
        except Exception as error:  # одна битая строка не должна ронять весь перенос
            errors.append(f"строка {index} ({row.get('Номер заказа')}): {str(error).splitlines()[0]}")
            importer.orders.pop(str(row.get("Номер заказа") or ""), None)
        if index % 1000 == 0:
            print(f"  … {index - 1} строк", flush=True)
    importer.finish()

    print("\nИтог:")
    for key, value in importer.stats.items():
        print(f"  {key}: {value}")
    for warning in importer.warnings:
        print(f"  ! {warning}")
    if errors:
        print(f"\nОшибок: {len(errors)} (первые 20):")
        for error in errors[:20]:
            print("  ", error)
    if args.commit:
        db.commit()
        print(f"\nЗаписано в базу за {time.time() - started:.0f} с.")
    else:
        db.rollback()
        print(f"\nПробный прогон за {time.time() - started:.0f} с — ничего не сохранено. Для записи добавьте --commit.")
    db.close()
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
